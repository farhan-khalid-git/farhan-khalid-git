"""Generate fictional example input files recognized by the workload importer.

Run from the project root: python examples/generate_input_templates.py
Requires openpyxl, installed through requirements.txt.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DIR = Path(__file__).resolve().parent


def save(rows, path, sheet_name):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='184C48')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for col in ws.columns:
        width = min(40, max(15, max(len(str(c.value or '')) for c in col) + 3))
        ws.column_dimensions[get_column_letter(col[0].column)].width = width
    ws.freeze_panes = 'A2'
    wb.save(path)
    print(f'Created: {path}')


if __name__ == '__main__':
    save([
        ['Name + CNIC','Course Code','Course Title','Credit Hours','Class','Semester','No. of Students'],
        ['Dr. Sample Teacher A','AGR-503','Crop Production Principles','3(2-1)','BS-AGES-BWP-5TH-AGRON-M','5th',25],
        [None,'AGR-505','Crop Physiology','3(2-1)','BS-AGES-BWP-5TH-AGRON-M','5th',25],
        ['Dr. Sample Teacher B','AGR-601','Field Crop Production','3(2-1)','BS-AGES-BWP-6TH-AGRON-M','6th',24],
        ['Dr. Sample Teacher C','AGR-701','Advanced Agronomic Research','3(3-0)','MSAGRO-1st-1E','1st',12],
        [None,'AGR-703','Experimental Techniques','3(2-1)','MSAGRO-1st-1E','1st',12],
        ['Dr. Sample Teacher D','AGR-705','Crop Climate Relationships','3(3+0)','MSAGRO-2nd-1E','2nd',10],
        ['Dr. Sample Teacher E','AGR-707','Field Sampling Lab','3(0-3)','MSAGRO-3rd-1E','3rd',8],
    ], DIR / 'sample_workload_annexure_a.xlsx', 'Teacher Wise Workload')
    save([
        ['Name of Teacher','Course Code','Course Title','Credit Hours','Associated Sections','Student Strength'],
        ['Dr. Sample Teacher A','AGR-503','Crop Production Principles','3(2-1)','BS-AGES-BWP-5TH-AGRON-M',25],
        [None,'AGR-505','Crop Physiology','3(2-1)','BS-AGES-BWP-5TH-AGRON-M',25],
        ['Dr. Sample Teacher B','AGR-601','Field Crop Production','3(2-1)','BS-AGES-BWP-6TH-AGRON-M',24],
        ['Dr. Sample Teacher C','AGR-701','Advanced Agronomic Research','3(3-0)','MSAGRO-1st-1E',12],
        [None,'AGR-703','Experimental Techniques','3(2-1)','MSAGRO-1st-1E',12],
        ['Dr. Sample Teacher D','AGR-705','Crop Climate Relationships','3(3+0)','MSAGRO-2nd-1E',10],
    ], DIR / 'sample_course_allocation.xlsx', 'Courses with MSc')
